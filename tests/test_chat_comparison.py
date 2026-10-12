import re
from io import BytesIO
from openpyxl import load_workbook
from app.main import app
from app.routes.web import ai_service
from app.schemas.proforma import ProformaDraft
from app.services.ai_service import AIResult
from app.services.chat_comparison import price_matrix
from tests.test_web import client


def test_chat_three_proformas_to_excel_without_request(client):
    calls=[]
    class Fake:
        def extract_comparison_block(self,text,filename,role,image=None):
            calls.append(text)
            return AIResult(ProformaDraft(supplier_name='Loja '+str(len(calls)),document_reference='PF',document_date=None,
                valid_until=None,warnings='',lines=[dict(description='PC i7 16GB',reference='ABC',quantity='2',unit='unidade',
                unit_price=str(100*len(calls)),currency='AOA',vat_included=True,evidence='Preço declarado no documento')]),'mock',10,10)
    app.dependency_overrides[ai_service]=Fake
    response=client.post('/chat',data={'message':'Compara estas proformas'},files=[('document',(f'pf{i}.txt',f'Loja {i}: PC i7'.encode(), 'text/plain')) for i in range(3)],headers={'HX-Request':'true'})
    assert response.status_code==200
    assert len(calls)==3
    assert 'Loja 1' in response.text and 'Loja 3' in response.text
    assert '100,00 Kz' in response.text and '300,00 Kz' in response.text
    url=re.search(r'href="(/comparisons/[^\"]+/export.xlsx)"',response.text).group(1)
    exported=client.get(url)
    assert exported.status_code==200
    book=load_workbook(BytesIO(exported.content))
    assert list(book['Comparação'].values)[4][-5:]==(100,200,300,100,300)
    assert book.sheetnames==['Comparação','Fontes']
    assert client.post('/chat',data={'mode':'compare'}).status_code==422
    assert client.get('/comparisons/00000000000000000000000000000000/export.xlsx').status_code==404


def test_matrix_does_not_mix_currencies_specs_or_duplicate_lines():
    line=dict(description='PC 16GB',reference='A',unit='unidade',currency='AOA',vat_included=None,unit_price='100',quantity='1')
    docs=[{'draft':{'lines':[line]}},{'draft':{'lines':[line|{'unit_price':'200'},line|{'description':'PC 8GB','unit_price':'50'}]}}]
    rows=price_matrix(docs)
    assert len(rows)==2 and rows[0]['lowest']=='100' and rows[0]['highest']=='200'
    docs[1]['draft']['lines'][0]['currency']='USD'
    assert all(row['lowest'] is None for row in price_matrix(docs))
    docs[0]['draft']['lines'].append(line)
    assert all(row['lowest'] is None for row in price_matrix(docs))


def test_wording_groups_equivalent_descriptions_but_preserves_specs():
    base=dict(reference='REF123',quantity='2',unit='unidade',unit_price='100',currency='AOA',vat_included=True)
    def compare(a,b):
        return price_matrix([{'draft':{'lines':[base|{'description':a}]}},
                             {'draft':{'lines':[base|{'description':b,'unit':'un','unit_price':'200'}]}}])
    rows=compare('Portátil HP i7 RAM 16 GB SSD 512 GB','Laptop HP i7 com 16GB RAM e SSD 512GB')
    assert len(rows)==1
    assert len(compare('Portátil HP i7 RAM 16 GB SSD 512 GB','Laptop HP i7 16GB RAM 512GB SSD'))==1
    assert len(compare('Portátil HP i7 RAM 16GB SSD 512GB','Notebook HP i7 RAM 8GB SSD 512GB'))==2
    assert len(compare('Portátil RAM 8GB SSD 16GB','Laptop RAM 16GB SSD 8GB'))==2


def test_chat_web_search_returns_comparison_and_excel(client, monkeypatch):
    from datetime import timedelta
    from types import SimpleNamespace
    from sqlalchemy import select
    from app.models import Supplier, PurchaseRequestItem, PriceSearch, SupplierOffer, utc_now
    from app.schemas.chat import KotaResponse, PurchaseDraft, PurchaseItemDraft
    import app.routes.web as web
    import app.routes.chat_search as search_route
    monkeypatch.setattr(search_route,'SessionLocal',web.SessionLocal)
    with web.SessionLocal() as session:
        suppliers=[Supplier(name=f'Loja {i}',website=f'https://example{i}.com',active=True,scraping_enabled=True) for i in range(2)]
        session.add_all(suppliers);session.commit();ids=[s.id for s in suppliers]
    class Fake:
        def respond(self,message,context):
            return AIResult(KotaResponse(reply='Rascunho',purchase_request=PurchaseDraft(title='PC',conditions=[],items=[PurchaseItemDraft(description='HP PC',normalized_description='HP PC',brand='HP',model=None,reference=None,quantity=2,unit='unidade',specifications=[])])),'mock',10,10)
    def start(request_id,item_ids,supplier_ids,refresh,factory):
        with factory() as session:
            search=PriceSearch(purchase_request_item_id=item_ids[0],search_query='HP',expires_at=utc_now()+timedelta(days=1),status='completed',raw_result_json={'supplier_ids':supplier_ids,'sources':[]})
            session.add(search);session.flush()
            for i,sid in enumerate(supplier_ids):
                session.add(SupplierOffer(purchase_request_item_id=item_ids[0],supplier_id=sid,price_search_id=search.id,product_name='HP PC',unit_price=100+i*100,currency='AOA',vat_included=True,match_confidence=95,evidence_json={}))
            session.commit()
    monkeypatch.setattr(search_route,'search_jobs',SimpleNamespace(start=start,status=lambda _: {'state':'completed','done':1,'total':1}))
    app.dependency_overrides[ai_service]=Fake
    result=client.post('/chat',data={'mode':'search','message':'Pesquisa 2 HP PC','supplier_ids':ids},headers={'HX-Request':'true'})
    assert result.status_code==200 and 'data-chat-search' in result.text
    url=re.search(r'data-chat-search="([^"]+)"',result.text).group(1)
    report=client.get(url)
    assert report.status_code==200 and '100,00 Kz' in report.json()['html']
    excel=re.search(r'href="(/comparisons/[^"]+/export.xlsx)"',report.json()['html']).group(1)
    assert client.get(excel).status_code==200
    assert client.post('/chat',data={'mode':'search','message':'Pesquisa HP'}).status_code==422
