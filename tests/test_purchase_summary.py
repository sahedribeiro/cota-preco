from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.models import PurchaseRequest, PurchaseRequestItem, SupplierOffer, utc_now
from app.services.purchase_summary import purchase_summary
from app.services.purchase_export import export_purchase
from tests.test_requests import client, save


def test_totals_separate_currency_and_vat_and_keep_missing_unknown():
    items, results = [], {}
    for i, (currency, vat, quantity) in enumerate([('AOA', True, '2'), ('AOA', False, '1'), ('USD', True, '3'), ('AOA', True, None)], 1):
        item = PurchaseRequestItem(id=i, description='PC', quantity=Decimal(quantity) if quantity else None)
        offer = SupplierOffer(supplier_id=1, unit_price=Decimal('10.25'), currency=currency, vat_included=vat,
                              match_confidence=95, evidence_json={})
        items.append(item)
        results[i] = dict(winner=offer, stale=False, names={1:'Loja'}, search=SimpleNamespace(status='completed'), reason='Teste')
    summary = purchase_summary(items, results)
    assert summary['totals'] == {('AOA',True):Decimal('20.50'), ('AOA',False):Decimal('10.25'), ('USD',True):Decimal('30.75')}
    assert summary['counted'] == 3 and summary['pending'] == 1
    results[1]['stale'] = True
    results[2]['search'].status = 'partial'
    summary = purchase_summary(items, results)
    assert ('AOA', True) not in summary['totals'] and summary['partial']


def test_export_is_readable_and_does_not_execute_untrusted_text(client):
    response = save(client)
    url = response.headers['location']
    client.post(url, data={'title':'=HYPERLINK("https://example.com")'})
    exported = client.get(url + '/export.xlsx')
    assert exported.status_code == 200
    assert 'attachment' in exported.headers['content-disposition']
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook['Resumo']
    assert sheet.freeze_panes and sheet.auto_filter.ref
    assert not any(cell.data_type == 'f' for row in sheet for cell in row)
    assert any(cell.value == 'Sem total calculável' for row in sheet for cell in row)
    assert any(cell.value == 'Ainda não pesquisado.' for row in sheet for cell in row)
    assert 'Resumo da compra' in client.get(url).text
    assert client.get('/requests/999999/export.xlsx').status_code == 404


def test_export_values_proformas_and_formula_injection():
    item = PurchaseRequestItem(id=1, description='=1+1', quantity=Decimal('2.5'), unit='unidade')
    offer = SupplierOffer(supplier_id=2, product_url='https://example.com/product', unit_price=Decimal('10.25'),
                          currency='AOA', vat_included=True, match_confidence=95, evidence_json={}, observed_at=utc_now())
    results = {1:dict(winner=offer, stale=False, names={2:'=DANGEROUS()'}, search=SimpleNamespace(status='completed'), reason='Fonte parcial')}
    summary = purchase_summary([item], results)
    assert summary['totals'][('AOA', True)] == Decimal('25.63')
    document = SupplierOffer(purchase_request_item_id=1, supplier_id=2, product_name='Proforma', unit_price=Decimal('999'),
                             currency='AOA', evidence_json={'document_reference':'=DANGEROUS()'})
    data = export_purchase(PurchaseRequest(id=1, title='Compra'), summary, [document], {2:'Loja'})
    book = load_workbook(BytesIO(data))
    assert book.sheetnames == ['Resumo', 'Proformas']
    assert not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row)
    values = [cell.value for row in book['Resumo'] for cell in row]
    assert 25.63 in values and 999 not in values and '=1+1' in values
    assert any(cell.value == 999 for row in book['Proformas'] for cell in row)
