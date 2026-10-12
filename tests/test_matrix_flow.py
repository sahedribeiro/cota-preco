import json
import re
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import shutil

import pymupdf
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func

from app.main import app
from app.routes import web
from app.models import PriceRecord
from app.services.comparison_documents import extract_document, table_text, request_table, aligned_proforma
from app.services.procurement_matrix import prepare_source, build_matrix
from app.services.matrix_quote import calculate_quote
from app.services.ai_service import AIResult
from app.schemas.comparison_document import ComparisonDocument
from tests.test_web import client

FIXTURE = Path(__file__).parent / 'fixtures' / 'girassol'
private_documents = pytest.mark.skipif(
    not all((FIXTURE / name).is_file() for name in ('oc.xlsx', 'neomed.pdf', 'AppyPay.pdf', 'mda.pdf', 'neomed.pdf.txt')),
    reason='Documentos comerciais privados disponíveis apenas na validação local',
)


class NoAI:
    def extract_comparison_block(self, *args):
        raise AssertionError('This supported local layout must not spend AI tokens')


def fixture_report(tmp_path):
    for file in FIXTURE.glob('*'):
        if file.suffix in {'.pdf', '.xlsx'}:
            shutil.copy(file, tmp_path / file.name)
    requested = extract_document(tmp_path/'oc.xlsx', 'oc.xlsx', 'request', NoAI(), lambda _: None)
    sources = [prepare_source(extract_document(tmp_path/name, name, 'supplier', NoAI(), lambda _: None))
               for name in ('neomed.pdf', 'AppyPay.pdf', 'mda.pdf')]
    return build_matrix(requested['lines'], sources)


@private_documents
def test_girassol_local_extraction_matrix_quote_and_alerts(tmp_path):
    report = fixture_report(tmp_path)
    assert len(report['rows']) == 37
    assert report['summary']['three_prices'] == 37
    assert report['summary']['gaps'] == report['summary']['review'] == 0
    requested = request_table(table_text(FIXTURE/'oc.xlsx'))['lines']
    assert [(row['number'], row['quantity']) for row in report['rows']] == [(row['number'], row['quantity']) for row in requested]
    names = [doc['label'] for doc in report['documents']]
    assert names[:2] == ['NEOMED', 'APPYPAY']
    assert names[2].startswith('MDA Saúde')
    assert all('Girassol' not in name and 'identificado' not in name for name in names)
    for row in report['rows']:
        assert Decimal(row['lowest']) == Decimal(row['cells'][2]['price'])
        assert Decimal(row['highest']) == Decimal(row['cells'][0]['price'])
    for actual, expected in zip(report['totals'], ['72905111.08', '64935726.47', '59032478.61']):
        assert abs(Decimal(actual)-Decimal(expected)) <= 5
    quote = calculate_quote(report, dict(strategy='highest', markup='60', vat='14'))
    assert [quote[k] for k in ('subtotal', 'tax', 'total', 'margin')] == ['116648184.64', '16330745.85', '132978930.49', '43743070.28']
    assert any('NEOMED' in alert and '1.235' in alert and '37/37' in alert for alert in report['alerts'])
    assert any('APPYPAY' in alert and '1.100' in alert and '37/37' in alert for alert in report['alerts'])
    # Printed unit price for 400 gowns differs from its printed line total by 2 Kz.
    assert report['documents'][0]['line_mismatches'] == ['13']


@private_documents
def test_end_to_end_import_idempotence_excel_and_customer_pdf(client):
    app.dependency_overrides[web.ai_service] = NoAI
    def submit():
        files = [('solicitation', ('oc.xlsx', (FIXTURE/'oc.xlsx').read_bytes()))]
        files += [('proformas', (name, (FIXTURE/name).read_bytes())) for name in ('neomed.pdf', 'AppyPay.pdf', 'mda.pdf')]
        return client.post('/matrices', data={'message': 'mais alto +60% + IVA, em nome da Empresa Teste', 'use_proformas': 'true'}, files=files, follow_redirects=False)
    response = submit()
    assert response.status_code == 303, response.text
    url = response.headers['location']
    page = client.get(url)
    assert page.status_code == 200, page.text
    assert 'APPYPAY' in page.text and '37/37' in page.text
    assert 'matrix-low' in page.text
    with web.SessionLocal() as session:
        assert session.scalar(select(func.count()).select_from(PriceRecord)) == 111
    assert submit().status_code == 303
    with web.SessionLocal() as session:
        assert session.scalar(select(func.count()).select_from(PriceRecord)) == 111
    catalog_only = client.post('/matrices', data={'use_catalog': 'true', 'catalog_days': '0'},
        files={'solicitation': ('oc.xlsx', (FIXTURE/'oc.xlsx').read_bytes())}, follow_redirects=False)
    assert catalog_only.status_code == 303, catalog_only.text
    catalog_page = client.get(catalog_only.headers['location'])
    assert 'Catálogo ·' in catalog_page.text and '37/37' in catalog_page.text
    logo = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    logo.clear_with(255)
    quote = client.post(url+'/quote', data={'strategy': 'highest', 'markup': '60', 'vat': '14', 'company': 'Empresa Teste', 'customer': 'Clínica Teste'},
                        files={'logo': ('logo.png', logo.tobytes('png'), 'image/png')}, follow_redirects=False)
    assert quote.status_code == 303, quote.text
    excel = client.get(url+'/export.xlsx')
    book = load_workbook(BytesIO(excel.content))
    assert book.sheetnames == ['Matriz', 'Cotação', 'Margem', 'Fontes']
    assert book['Matriz'].max_row == 39
    assert book['Matriz']['G2'].fill.fgColor.rgb.endswith('DCFCE7')
    assert book['Matriz']['E2'].fill.fgColor.rgb.endswith('FEE2E2')
    assert book['Fontes'].max_row == 112
    pdf = client.get(url+'/quote.pdf')
    assert pdf.status_code == 200
    with pymupdf.open(stream=pdf.content, filetype='pdf') as document:
        assert document[0].get_images()
        text = ''.join(page.get_text() for page in document)
        assert 'NEOMED' not in text and 'APPYPAY' not in text and 'Margem' not in text and 'Custo' not in text
        assert '132.978.930,49' in text or '132 978 930,49' in text
        assert 'Empresa Teste' in text
    assert client.get(url+'/matrix.pdf').status_code == 200


@private_documents
def test_vat_unknown_compares_but_incompatible_vat_and_currencies_do_not(tmp_path):
    report = fixture_report(tmp_path)
    item = report['items'][0]
    sources = report['documents'][:2]
    for source in sources:
        source['draft']['lines'] = source['draft']['lines'][:1]
    assert build_matrix([item], sources)['rows'][0]['lowest']
    sources[0]['draft']['lines'][0]['vat_included'] = True
    sources[1]['draft']['lines'][0]['vat_included'] = False
    assert build_matrix([item], sources)['rows'][0]['lowest'] is None
    sources[1]['draft']['lines'][0]['vat_rate'] = '14'
    assert build_matrix([item], sources)['rows'][0]['lowest']
    sources[1]['draft']['lines'][0]['currency'] = 'USD'
    assert build_matrix([item], sources)['rows'][0]['lowest'] is None


@private_documents
def test_ai_blocks_are_checkpointed_and_merge_all_lines(tmp_path):
    text = (FIXTURE/'neomed.pdf.txt').read_text(encoding='utf-8')
    draft = aligned_proforma(text)
    path = tmp_path/'unknown-layout.txt'
    path.write_text('\n'.join('ARTIGO '+str(i)+' '+('x'*130) for i in range(180)), encoding='utf-8')
    calls = []
    class RecordedAI:
        def extract_comparison_block(self, text, filename, role, image):
            index = len(calls)
            calls.append(text)
            data = dict(draft, lines=draft['lines'][index*7:(index+1)*7])
            return AIResult(ComparisonDocument.model_validate(data), 'recorded-local-fixture', 0, 0)
    result = extract_document(path, path.name, 'supplier', RecordedAI(), lambda _: None)
    assert len(result['lines']) == 37
    assert len(calls) > 5
    extract_document(path, path.name, 'supplier', NoAI(), lambda _: None)


def test_empty_solicitation_and_six_proformas_rejected(client):
    assert client.post('/matrices').status_code == 422
    files = [('proformas', (f'{i}.txt', b'test')) for i in range(6)]
    assert client.post('/matrices', data={'message': 'PC'}, files=files).status_code == 422


@private_documents
def test_preview_and_duplicate_attachments_keep_one_supplier(client):
    app.dependency_overrides[web.ai_service] = NoAI
    content = (FIXTURE/'neomed.pdf').read_bytes()
    preview = client.post('/matrices/inspect', files={'proformas': ('neomed.pdf', content)})
    assert preview.status_code == 200 and 'NEOMED' in preview.text and '37 artigos' in preview.text
    files = [('solicitation', ('oc.xlsx', (FIXTURE/'oc.xlsx').read_bytes())),
             ('proformas', ('neomed.pdf', content)), ('proformas', ('neomed.pdf', content))]
    response = client.post('/matrices', data={'use_proformas': 'true'}, files=files, follow_redirects=False)
    assert response.status_code == 303
    rid = response.headers['location'].split('/')[-1]
    report = json.loads((web.settings.upload_path/f'{rid}.comparison.json').read_text(encoding='utf-8'))
    assert report['summary']['suppliers'] == 1
    assert all(row['count'] == 1 for row in report['rows'])
    with web.SessionLocal() as session:
        assert session.scalar(select(func.count()).select_from(PriceRecord)) == 37


@private_documents
def test_two_prices_same_currency_still_compare_with_foreign_third(tmp_path):
    report = fixture_report(tmp_path)
    report['documents'][0]['draft']['lines'][0]['currency'] = 'USD'
    matrix = build_matrix(report['items'], report['documents'])
    row = matrix['rows'][0]
    assert row['currency'] == 'AOA'
    assert row['lowest'] == row['cells'][2]['price']
    assert row['highest'] == row['cells'][1]['price']
    assert matrix['totals'][0] is None


@private_documents
def test_low_confidence_blocks_issue_and_manual_and_exemption_work(tmp_path):
    report = fixture_report(tmp_path)
    options = dict(strategy='manual', markup='10', vat='0', exemption='Menção fornecida pelo utilizador',
                   manual={str(index): '2' for index in range(37)})
    assert calculate_quote(report, options)['tax'] == '0.00'
    with pytest.raises(ValueError, match='menção legal'):
        calculate_quote(report, options | {'exemption': ''})
    report['rows'][0]['cells'][2]['confidence'] = 70
    with pytest.raises(ValueError, match='confirma o preço'):
        calculate_quote(report, options)


def test_ai_matching_uses_candidates_and_keeps_conflicting_specs_separate():
    from app.services.procurement_matrix import assist_matching
    from app.schemas.comparison_document import ItemMatches
    item = dict(description='Luvas latex tamanho M caixa 100', quantity='2', unit='un')
    line = dict(description='Luvas em latex tam M caixa 100', quantity='2', unit='un', unit_price='100', currency='AOA')
    sources = [dict(label='Loja', draft={'lines': [line]})]
    calls = []
    class FakeMatcher:
        def match_comparison_lines(self, candidates):
            calls.append(candidates)
            return AIResult(ItemMatches(matches=[dict(item_index=0, line_index=0, confidence=95, reason='Mesmo produto e embalagem')]), 'mock', 1, 1)
    assist_matching([item], sources, FakeMatcher(), lambda _: None)
    assert len(calls) == 1 and len(calls[0][0]['candidates']) == 1
    assert build_matrix([item], sources)['rows'][0]['count'] == 1
    line['description'] = 'Luvas latex tamanho M caixa 50'
    sources = [dict(label='Loja', draft={'lines': [line]})]
    assist_matching([item], sources, FakeMatcher(), lambda _: None)
    assert len(calls) == 1
    assert build_matrix([item], sources)['rows'][0]['count'] == 0
