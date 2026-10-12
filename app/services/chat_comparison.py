"""Side-by-side documentary prices, separate from purchase approval."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from app.services.comparison_service import normalize
from app.services.product_wording import canonical_description, canonical_unit


def price_matrix(documents: list[dict]) -> list[dict]:
    groups: dict[tuple, dict] = {}
    for column, document in enumerate(documents):
        for line in document['draft']['lines']:
            key = (canonical_description(line['description']), canonical_unit(line.get('unit')),
                   normalize(line.get('reference') or ''), line.get('currency'), line.get('vat_included'))
            row = groups.setdefault(key, dict(product=line['description'], unit=line.get('unit'),
                reference=line.get('reference'), currency=line.get('currency'), vat=line.get('vat_included'),
                offers=[[] for _ in documents]))
            try:
                amount = Decimal(line.get('unit_price') or '')
                price = str(amount) if amount.is_finite() and 0 < amount < Decimal('1e16') else None
            except InvalidOperation:
                price = None
            row['offers'][column].append({'price':price,'quantity':line.get('quantity'),'evidence':line.get('evidence',''),'description':line['description']})
    for row in groups.values():
        values = [Decimal(offers[0]['price']) for offers in row['offers'] if len(offers)==1 and offers[0]['price']]
        comparable = len(values) >= 2 and all(len(offers)<=1 for offers in row['offers']) and bool(row['currency'] and row['unit'])
        row['lowest'] = str(min(values)) if comparable else None
        row['highest'] = str(max(values)) if comparable else None
    return list(groups.values())


def comparison_excel(report: dict) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = 'Comparação'
    sheet.append(['Comparação de preços unitários declarados — não constitui recomendação de compra.'])
    sheet.append(['Associação conservadora por descrição normalizada, referência e unidade equivalentes. Moedas e IVA separados; IVA desconhecido = comparação nominal.'])
    sheet.append(['Pedido', report.get('message','')])
    sheet.append(['Produto','Referência','Unidade','Moeda','IVA'] + [d['label'] for d in report['documents']] + ['Menor preço','Maior preço'])
    for row in report['rows']:
        values = []
        for offers in row['offers']:
            values.append(Decimal(offers[0]['price']) if len(offers)==1 and offers[0]['price'] else 'Várias linhas: ver Fontes' if len(offers)>1 else 'Sem preço')
        sheet.append([row['product'],row['reference'],row['unit'],row['currency'],
                      'Incluído' if row['vat'] is True else 'Excluído' if row['vat'] is False else 'Desconhecido'] + values +
                     [Decimal(row['lowest']) if row['lowest'] else None,Decimal(row['highest']) if row['highest'] else None])
        for cell in list(sheet[sheet.max_row])[5:]:
            cell.number_format = '#,##0.00'
    sheet.freeze_panes = 'F5'
    sheet.auto_filter.ref = f'A4:{sheet.cell(4,sheet.max_column).column_letter}{sheet.max_row}'
    source = book.create_sheet('Fontes')
    source.append(['Fornecedor / ficheiro','Documento','Data declarada','Validade declarada','Avisos','Produto','Quantidade','Unidade','Preço original','Moeda','Evidência'])
    for document in report['documents']:
        draft=document['draft']
        for line in draft['lines']:
            source.append([document['label'],draft.get('document_reference'),draft.get('document_date'),draft.get('valid_until'),draft.get('warnings'),line['description'],line.get('quantity'),line.get('unit'),line.get('unit_price'),line.get('currency'),line.get('evidence')])
    for tab in book:
        for cells in tab:
            for cell in cells:
                if isinstance(cell.value,str):
                    cell.value=ILLEGAL_CHARACTERS_RE.sub('',cell.value)
                    cell.data_type='s'
                elif isinstance(cell.value,Decimal) and len(cell.value.as_tuple().digits)>15:
                    cell.value=str(cell.value); cell.data_type='s'
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width=28
        for cell in tab[4 if tab==sheet else 1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='16303B')
    output=BytesIO(); book.save(output)
    return output.getvalue()
