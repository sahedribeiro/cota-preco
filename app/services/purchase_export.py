from math import ceil
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.models import PurchaseRequest, SupplierOffer, utc_now


def export_purchase(purchase: PurchaseRequest, summary: dict, documents: list[SupplierOffer], names: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Resumo'
    sheet.append([f'Cota.Preço.AI · Requisição #{purchase.id} · {purchase.title}'])
    sheet.merge_cells('A1:N1')
    sheet.append(['Exportado (UTC)', utc_now().strftime('%d/%m/%Y %H:%M')])
    sheet.append(['Artigos com total conhecido', summary['counted'], 'Artigos pendentes', summary['pending']])
    sheet.append(['Comparação parcial de fontes', 'Sim' if summary['partial'] else 'Não'])
    sheet.append(['Totais conhecidos; não constituem orçamento final. Transporte e impostos em falta não calculados.'])
    sheet.merge_cells('A5:N5')
    sheet.row_dimensions[5].height = 30
    sheet.append(['Moeda', 'IVA', 'Subtotal conhecido'])
    for (currency, vat), total in sorted(summary['totals'].items()):
        sheet.append([currency, 'Incluído' if vat else 'Excluído', total])
        sheet.cell(sheet.max_row, 3).number_format = '#,##0.00'
    if not summary['totals']:
        sheet.append(['Sem total calculável'])
    sheet.append([])
    header = sheet.max_row + 1
    sheet.append(['Artigo', 'Produto pedido', 'Fornecedor seleccionado', 'Preço unitário', 'Moeda', 'Quantidade',
                  'Unidade', 'Total conhecido', 'IVA', 'Confiança automática', 'Fonte', 'Observado (UTC)', 'Pendências / limites'])
    for row in summary['rows']:
        item, offer = row['item'], row['offer']
        sheet.append([item.id, item.description, row['supplier'], offer.unit_price if offer else None,
                      offer.currency if offer else None, item.quantity, item.unit, row['total'],
                      ('Incluído' if row['vat'] else 'Excluído') if offer else None,
                      offer.match_confidence if offer else None, (('Proforma: ' + str(offer.evidence_json.get('document_reference') or offer.evidence_json.get('import_id'))) if offer.document_line_key else offer.product_url) if offer else None,
                      offer.observed_at.strftime('%d/%m/%Y %H:%M') if offer else None, row['reason'], row['origin']])
        for col in (4, 8):
            sheet.cell(sheet.max_row, col).number_format = '#,##0.00'
        sheet.cell(sheet.max_row, 6).number_format = '#,##0.###'
    sheet.freeze_panes = f'C{header + 1}'
    sheet.auto_filter.ref = f'A{header}:N{sheet.max_row}'
    if documents:
        doc = workbook.create_sheet('Proformas')
        doc.append(['Ofertas documentais — excluídas dos totais; validade, equivalência e stock por confirmar.'])
        doc.merge_cells('A1:J1')
        doc.append(['Artigo', 'Produto cotado', 'Fornecedor', 'Preço unitário', 'Moeda', 'IVA', 'Data declarada', 'Validade declarada', 'Documento', 'Revisor'])
        for offer in documents:
            e = offer.evidence_json
            doc.append([offer.purchase_request_item_id, offer.product_name, names[offer.supplier_id], offer.unit_price,
                        offer.currency, 'Incluído' if offer.vat_included is True else 'Excluído' if offer.vat_included is False else 'Desconhecido',
                        e.get('document_date'), e.get('valid_until'), e.get('document_reference'), e.get('reviewer')])
            doc.cell(doc.max_row, 4).number_format = '#,##0.00'
        doc.freeze_panes = 'C3'
        doc.auto_filter.ref = f'A2:J{doc.max_row}'
    for tab in workbook:
        tab.sheet_view.showGridLines = False
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub('', cell.value)
                    cell.data_type = 's'  # Untrusted descriptions must never become Excel formulas.
                elif isinstance(cell.value, Decimal) and len(cell.value.as_tuple().digits) > 15:
                    cell.value = str(cell.value)  # Preserve values beyond Excel's numeric precision.
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Calibri', size=11)
        for cell in tab[1]:
            cell.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='16303B')
        heading_row = header if tab.title == 'Resumo' else 2
        for cell in tab[heading_row]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='315A65')
        for col in range(1, tab.max_column + 1):
            tab.column_dimensions[get_column_letter(col)].width = 18 if col not in (2, 3, 11, 13) else 44
        tab.row_dimensions[1].height = 36
        for row in range(heading_row + 1, tab.max_row + 1):
            lines = max(ceil(len(str(cell.value or '')) / (tab.column_dimensions[cell.column_letter].width * .9)) for cell in tab[row])
            tab.row_dimensions[row].height = max(32, lines * 16)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
