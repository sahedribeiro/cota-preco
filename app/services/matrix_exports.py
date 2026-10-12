from decimal import Decimal
from io import BytesIO

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from app.core.formatting import money
from app.core.config import settings


def matrix_excel(report: dict) -> bytes:
    book = Workbook()
    matrix = book.active
    matrix.title = 'Matriz'
    matrix.append(['Nº', 'Item', 'Qtd', 'Un'] + [doc['label'] for doc in report['documents']] + ['Mais baixo', 'Mais alto', 'Dif. %'])
    for row in report['rows']:
        prices = [Decimal(cell['compared_price']) if cell and cell.get('compared_price') else None for cell in row['cells']]
        matrix.append([row['number'], row['product'], Decimal(row['quantity']) if row['quantity'] else None, row['unit']] + prices +
            [Decimal(row[key]) if row.get(key) is not None else None for key in ('lowest', 'highest', 'difference_percent')])
        for index, price in enumerate(prices, 5):
            if price is None:
                continue
            if row['lowest'] and price == Decimal(row['lowest']):
                matrix.cell(matrix.max_row, index).fill = PatternFill('solid', fgColor='DCFCE7')
            elif row['highest'] and price == Decimal(row['highest']):
                matrix.cell(matrix.max_row, index).fill = PatternFill('solid', fgColor='FEE2E2')
    matrix.append(['', 'TOTAL', '', ''] + [Decimal(value) if value is not None else None for value in report['totals']] +
        [Decimal(report['summary'][key]) if report['summary'][key] is not None else None for key in ('lowest', 'highest')])
    quote_tab = book.create_sheet('Cotação')
    margin_tab = book.create_sheet('Margem')
    quote = report.get('quote')
    if quote:
        quote_tab.append(['Empresa', quote['options'].get('company', '')])
        quote_tab.append(['Cliente', quote['options'].get('customer', '')])
        quote_tab.append(['Nº', 'Item', 'Qtd', 'Un', 'Preço venda', 'Total'])
        margin_tab.append(['Nº', 'Item', 'Custo unitário', 'Margem bruta'])
        for line in quote['lines']:
            quote_tab.append([line['number'], line['description'], Decimal(line['quantity']), line['unit'], Decimal(line['sale_price']), Decimal(line['total'])])
            margin_tab.append([line['number'], line['description'], Decimal(line['cost']), Decimal(line['margin'])])
        for label, key in [('Subtotal', 'subtotal'), ('IVA', 'tax'), ('TOTAL', 'total')]:
            quote_tab.append([label, Decimal(quote[key])])
        margin_tab.append(['TOTAL', '', '', Decimal(quote['margin'])])
    else:
        quote_tab.append(['Escolhe e guarda a estratégia no ecrã para incluir a cotação.'])
        margin_tab.append(['Sem cotação calculada. Uso interno.'])
    sources = book.create_sheet('Fontes')
    sources.append(['Fornecedor', 'Ficheiro', 'Referência', 'Data', 'Item', 'Descrição original', 'Qtd cotada', 'Preço original', 'Moeda', 'AOA assumida', 'IVA', 'Evidência', 'Origem'])
    for index, doc in enumerate(report['documents']):
        for row in report['rows']:
            cell = row['cells'][index]
            if cell:
                sources.append([doc['label'], doc['draft'].get('filename'), doc['draft'].get('document_reference'), cell.get('source_date'),
                    row['number'], cell['description'], cell.get('quantity'), cell.get('price'), cell.get('currency'), bool(cell.get('currency_assumed')),
                    'Incluído' if cell.get('vat_included') is True else 'Excluído' if cell.get('vat_included') is False else 'IVA não indicado', cell.get('evidence'), cell.get('source_kind')])
    for tab in book:
        tab.freeze_panes = 'E2' if tab == matrix else 'A2'
        tab.auto_filter.ref = tab.dimensions
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub('', cell.value)
                    cell.data_type = 's'
                elif isinstance(cell.value, (Decimal, int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = '#,##0.00'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in tab[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='16303B')
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = 24
        tab.column_dimensions['B'].width = 55
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def matrix_pdf(report: dict, customer: bool = False) -> bytes:
    quote = report.get('quote')
    if customer and not quote:
        raise ValueError('Escolhe e guarda uma estratégia antes de gerar a cotação.')
    if customer:
        options = quote['options']
        if not options.get('company', '').strip() or not options.get('customer', '').strip():
            raise ValueError('Preenche empresa emissora e cliente.')
        title = 'Cotação'
        headers = ['Nº', 'Item', 'Qtd', 'Un', 'Preço unitário', 'Total']
        rows = [[line['number'], line['description'], line['quantity'], line['unit'], money(line['sale_price'], quote['currency']), money(line['total'], quote['currency'])] for line in quote['lines']]
    else:
        title = 'Matriz comparativa · Uso interno'
        headers = ['Nº', 'Item', 'Qtd', 'Un'] + [doc['label'] for doc in report['documents']] + ['Mais baixo', 'Mais alto']
        rows = [[row['number'], row['product'], row['quantity'], row['unit']] +
                [money(cell['compared_price'], cell['currency']) if cell and cell.get('compared_price') else '—' for cell in row['cells']] +
                [money(row[key], row['currency']) if row.get(key) else '—' for key in ('lowest', 'highest')] for row in report['rows']]
        rows.append(['', 'TOTAL', '', ''] + [money(value, report['summary']['currency']) if value else '—' for value in report['totals']] +
                    [money(report['summary'][key], report['summary']['currency']) if report['summary'][key] else '—' for key in ('lowest', 'highest')])
    pdf = pymupdf.open()
    page_rect = pymupdf.paper_rect('a4' if customer else 'a3-l')
    margin = 28
    width = page_rect.width - 2 * margin
    widths = ([25, width-250, 35, 30, 75, 85] if customer else
              [28, 260, 38, 35] + [(width-361)/(len(headers)-4)]*(len(headers)-4))
    font_size, line_height, padding = 8, 11, 5
    page = pdf.new_page(width=page_rect.width, height=page_rect.height)
    y = 34

    def wrap(value: object, available: float, size: float = font_size) -> list[str]:
        text = str(value or '').replace('\n', ' ')
        result, line = [], ''
        for word in text.split():
            if pymupdf.get_text_length((line+' '+word).strip(), fontsize=size) <= available:
                line = (line+' '+word).strip()
                continue
            if line:
                result.append(line)
                line = ''
            for char in word:
                if pymupdf.get_text_length(line+char, fontsize=size) > available and line:
                    result.append(line)
                    line = ''
                line += char
        result.append(line)
        return result

    def paragraph(text: str, size: float = 9) -> None:
        nonlocal page, y
        for line in wrap(text, width-100 if y < 90 else width, size):
            if y > page_rect.height-48:
                page = pdf.new_page(width=page_rect.width, height=page_rect.height)
                y = 34
            page.insert_text((margin,y), line, fontsize=size, color=(.09,.19,.23))
            y += size+5
        y += 5

    paragraph(title, 17)
    if customer:
        paragraph(options.get('company',''), 13)
        for text in [f"NIF: {options.get('tax_id','')}", options.get('address',''), options.get('contacts',''), f"Cliente: {options.get('customer','')}"]:
            if text:
                paragraph(text)

    def draw_row(values: list, header: bool = False, data_index: int | None = None) -> None:
        nonlocal y, page
        wrapped = [wrap(value, column_width-2*padding) for value,column_width in zip(values,widths)]
        height = max(len(lines) for lines in wrapped)*line_height+2*padding
        if y+height > page_rect.height-40:
            page = pdf.new_page(width=page_rect.width, height=page_rect.height)
            y = 34
            if not header:
                draw_row(headers, True)
        x = margin
        for index,(lines,column_width) in enumerate(zip(wrapped,widths)):
            fill = (.91,.94,.95) if header else (1,1,1)
            if not customer and data_index is not None and data_index < len(report['rows']) and 4 <= index < 4+len(report['documents']):
                item = report['rows'][data_index]
                cell = item['cells'][index-4]
                if cell and cell.get('compared_price') and cell['currency'] == item['currency']:
                    price = Decimal(cell['compared_price'])
                    if item['lowest'] and price == Decimal(item['lowest']): fill = (.86,.98,.9)
                    elif item['highest'] and price == Decimal(item['highest']): fill = (1,.89,.89)
            page.draw_rect(pymupdf.Rect(x,y,x+column_width,y+height), color=(.8,.84,.86), fill=fill, width=.4)
            for li,line in enumerate(lines):
                page.insert_text((x+padding,y+padding+font_size+li*line_height),line,fontsize=font_size,fontname='hebo' if header else 'helv',color=(.09,.19,.23))
            x += column_width
        y += height

    draw_row(headers, True)
    for index,row in enumerate(rows):
        draw_row(row, data_index=index)
    y += 18
    if customer:
        for label,key in [('Subtotal','subtotal'),('IVA','tax'),('TOTAL','total')]:
            paragraph(label+': '+money(quote[key],quote['currency']), 11)
        for label,key in [('Validade','validity'),('Prazo','delivery'),('Pagamento','payment'),('Isenção','exemption')]:
            if options.get(key): paragraph(label+': '+options[key])
    else:
        paragraph('Preços unitários. Moeda assumida, IVA e evidência nas Fontes do Excel. Confirma condições comerciais.')
    if customer and quote['options'].get('logo'):
        import re
        logo = quote['options']['logo']
        if re.fullmatch(r'[a-f0-9]{32}\.png', logo) and (settings.upload_path / logo).is_file():
            pdf[0].insert_image(pymupdf.Rect(page_rect.width-115,24,page_rect.width-28,80),filename=str(settings.upload_path/logo))
    for index,page in enumerate(pdf):
        page.insert_text((margin,page_rect.height-20),f'{title} | {index+1}/{len(pdf)}',fontsize=8,color=(.4,.4,.4))
    output = pdf.tobytes()
    pdf.close()
    return output
