import csv
import io
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.config import settings
from app.services.file_storage_service import DocumentError


class VisualProcessingRequired(DocumentError):
    pass


def bounded_text(parts: list[str]) -> str:
    text = "\n".join(parts).strip()
    if len(text) > settings.document_max_chars:
        raise DocumentError(f"O documento ultrapassa {settings.document_max_chars} caracteres. Divide-o em partes menores; não enviei texto parcial à IA.")
    return text


class DocumentService:
    def extract(self, path: Path) -> str:
        try:
            if path.suffix == ".pdf":
                text = self._pdf(path)
            elif path.suffix == ".xlsx":
                text = self._xlsx(path)
            else:
                raw = path.read_bytes()
                try:
                    text = raw.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = raw.decode("cp1252")
                if "\x00" in text:
                    raise DocumentError("Usa um ficheiro de texto UTF-8 ou Windows-1252.")
                if path.suffix == ".csv":
                    text = self._csv(text)
            text = bounded_text([text])
            if not text.strip():
                raise DocumentError("Não encontrei texto ou células preenchidas no documento.")
            return text
        except DocumentError:
            raise
        except Exception:
            raise DocumentError("Não foi possível ler o documento. Confirma o formato e se o ficheiro está válido e sem palavra-passe.") from None

    def _csv(self, text: str) -> str:
        bounded_text([text])
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(io.StringIO(text), dialect=dialect, strict=True)
        return bounded_text([" | ".join(row) for row in rows])

    def _pdf(self, path: Path) -> str:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentError("PDF protegido. Anexa uma versão sem palavra-passe.")
        if len(reader.pages) > 30:
            raise DocumentError("O PDF excede 30 páginas. Divide-o em documentos menores.")
        parts: list[str] = []
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise VisualProcessingRequired("Este PDF tem páginas sem texto extraível. Ficou marcado para processamento visual futuro; OCR ainda não está disponível. Envia uma versão com texto, XLSX, CSV ou TXT.")
            parts.append(f"Página {number}\n{text}")
            bounded_text(parts)
        return bounded_text(parts)

    def _xlsx(self, path: Path) -> str:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 30 * 1024 * 1024:
                raise DocumentError("A folha de cálculo é demasiado grande depois de descomprimida.")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        parts: list[str] = []
        has_values = False
        try:
            for sheet in workbook:
                if (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 40:
                    raise DocumentError("Cada folha pode ter até 500 linhas e 40 colunas. Envia apenas a área da requisição.")
                parts.append(f"Folha: {sheet.title}")
                for row in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise DocumentError("A folha contém fórmulas. Exporta uma cópia apenas com valores para evitar interpretar resultados desactualizados.")
                    values = [str(cell.value) if cell.value is not None else "" for cell in row]
                    if any(values):
                        has_values = True
                        parts.append(" | ".join(values))
                        bounded_text(parts)
            return bounded_text(parts) if has_values else ""
        finally:
            workbook.close()
